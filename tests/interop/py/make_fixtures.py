"""Genera archivos de prueba con implementaciones independientes (python-docx, openpyxl, python-pptx)."""
import sys, os, struct, zlib
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt
from pptx.dml.color import RGBColor as PRGB
from pptx.enum.shapes import MSO_SHAPE

out = sys.argv[1]
os.makedirs(out, exist_ok=True)

def png(path, w=4, h=3):
    raw = b''.join(b'\x00' + b'\xff\x00\x00' * w for _ in range(h))
    def chunk(t, d):
        c = struct.pack('>I', len(d)) + t + d
        return c + struct.pack('>I', zlib.crc32(t + d) & 0xffffffff)
    data = b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(raw)) + chunk(b'IEND', b'')
    open(path, 'wb').write(data)

png(os.path.join(out, 'pixel.png'))

# ---------- DOCX ----------
d = Document()
d.add_heading('Informe de interoperabilidad', level=1)
p = d.add_paragraph('Texto normal con ')
r = p.add_run('negrita'); r.bold = True
p.add_run(' y ')
r = p.add_run('cursiva roja'); r.italic = True; r.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
p.add_run(' y ')
r = p.add_run('grande'); r.font.size = Pt(20)
c = d.add_paragraph('Centrado'); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
d.add_heading('Listas', level=2)
d.add_paragraph('Elemento uno', style='List Bullet')
d.add_paragraph('Elemento dos', style='List Bullet')
d.add_paragraph('Paso uno', style='List Number')
d.add_paragraph('Paso dos', style='List Number')
t = d.add_table(rows=2, cols=3)
t.style = 'Table Grid'
for i, txt in enumerate(['A', 'B', 'C']): t.cell(0, i).text = txt
for i, txt in enumerate(['1', '2', '3']): t.cell(1, i).text = txt
a = t.cell(0, 0).merge(t.cell(0, 1))
d.add_picture(os.path.join(out, 'pixel.png'), width=Inches(1))
d.add_page_break()
d.add_paragraph('Segunda página')
s = d.sections[0]
s.orientation = WD_ORIENT.LANDSCAPE
s.page_width, s.page_height = s.page_height, s.page_width
s.header.paragraphs[0].text = 'Cabecera de prueba'
d.save(os.path.join(out, 'python-docx.docx'))

# ---------- XLSX ----------
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Datos'
rows = [['Producto', 'Unidades', 'Precio'], ['Café', 10, 2.5], ['Té', 4, 3], ['Agua', 25, 0.8]]
for r in rows: ws.append(r)
formulas = {
    'D2': '=B2*C2', 'D3': '=B3*C3', 'D4': '=B4*C4', 'D5': '=SUM(D2:D4)',
    'E2': '=IF(D2>20,"alto","bajo")', 'E3': '=ROUND(AVERAGE(C2:C4),2)', 'E4': '=VLOOKUP("Té",A2:C4,3,FALSE)',
    'E5': '=SUMIF(B2:B4,">5",D2:D4)', 'F2': '=COUNTIF(A2:A4,"*a*")', 'F3': '=MAX(B2:B4)-MIN(B2:B4)',
    'F4': '=CONCATENATE(A2," y ",A3)', 'F5': '=Resumen!A1*2', 'G2': '=DATE(2025,3,15)', 'G3': '=YEAR(G2)+MONTH(G2)',
    'G4': '=INDEX(A2:C4,2,1)&MATCH(25,B2:B4,0)', 'G5': '=IFERROR(1/0,"div0")', 'H2': '=SUMPRODUCT(B2:B4,C2:C4)',
    'H3': '=LEFT(UPPER(A4),3)&LEN(A2)', 'H4': '=MOD(17,5)+POWER(2,3)+ABS(-1)', 'H5': '=AND(B2>1,OR(B3>10,B4>10))',
}
for k, v in formulas.items(): ws[k] = v
ws['A1'].font = Font(bold=True, color='FFFFFF')
ws['A1'].fill = PatternFill('solid', fgColor='0F766E')
ws['C2'].number_format = '#,##0.00'
ws['B1'].alignment = Alignment(horizontal='center')
ws.merge_cells('A7:C7')
ws['A7'] = 'Celda combinada'
ws.column_dimensions['A'].width = 20
ws.freeze_panes = 'A2'
ws2 = wb.create_sheet('Resumen')
ws2['A1'] = 21
wb.save(os.path.join(out, 'openpyxl.xlsx'))

# ---------- PPTX ----------
pr = Presentation()
s1 = pr.slides.add_slide(pr.slide_layouts[0])
s1.shapes.title.text = 'Título desde python-pptx'
s1.placeholders[1].text = 'Subtítulo'
s2 = pr.slides.add_slide(pr.slide_layouts[1])
s2.shapes.title.text = 'Puntos'
tf = s2.placeholders[1].text_frame
tf.text = 'Primer punto'
tf.add_paragraph().text = 'Segundo punto'
sh = s2.shapes.add_shape(MSO_SHAPE.OVAL, PInches(6), PInches(4), PInches(2), PInches(1.5))
sh.fill.solid(); sh.fill.fore_color.rgb = PRGB(0x22, 0x88, 0xEE)
sh.text = 'Óvalo'
s3 = pr.slides.add_slide(pr.slide_layouts[6])
s3.shapes.add_picture(os.path.join(out, 'pixel.png'), PInches(1), PInches(1), PInches(2))
tb = s3.shapes.add_textbox(PInches(4), PInches(1), PInches(4), PInches(1))
run = tb.text_frame.paragraphs[0].add_run(); run.text = 'Texto libre'; run.font.bold = True; run.font.size = PPt(32)
tbl = s3.shapes.add_table(2, 2, PInches(1), PInches(4), PInches(6), PInches(1.5)).table
tbl.cell(0, 0).text = 'H1'; tbl.cell(0, 1).text = 'H2'; tbl.cell(1, 0).text = 'v1'; tbl.cell(1, 1).text = 'v2'
s3.notes_slide.notes_text_frame.text = 'Notas desde python-pptx'
pr.save(os.path.join(out, 'python-pptx.pptx'))
print('ok')
