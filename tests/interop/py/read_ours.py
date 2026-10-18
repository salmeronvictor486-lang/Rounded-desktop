"""Lee con bibliotecas independientes los archivos exportados por Rounded Desktop y vuelca un resumen JSON."""
import sys, json
from docx import Document
import openpyxl
from pptx import Presentation

kind, path = sys.argv[1], sys.argv[2]
out = {}
if kind == 'docx':
    d = Document(path)
    out['paragraphs'] = [{'text': p.text, 'style': p.style.name, 'bold': [r.bold for r in p.runs if r.text]} for p in d.paragraphs]
    out['tables'] = [[[c.text for c in row.cells] for row in t.rows] for t in d.tables]
    s = d.sections[0]
    out['landscape'] = s.page_width > s.page_height
    out['header'] = s.header.paragraphs[0].text if s.header.paragraphs else ''
    out['images'] = len(d.inline_shapes)
elif kind == 'xlsx':
    wb = openpyxl.load_workbook(path)
    wbv = openpyxl.load_workbook(path, data_only=True)
    out['sheets'] = wb.sheetnames
    ws = wb[wb.sheetnames[0]]; wv = wbv[wb.sheetnames[0]]
    out['cells'] = {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}
    out['values'] = {c.coordinate: c.value for row in wv.iter_rows() for c in row if c.value is not None}
    out['bold_A1'] = bool(ws['A1'].font.bold)
    out['fill_A1'] = ws['A1'].fill.fgColor.rgb if ws['A1'].fill and ws['A1'].fill.fgColor else None
    out['merged'] = [str(m) for m in ws.merged_cells.ranges]
    out['freeze'] = ws.freeze_panes
elif kind == 'pptx':
    pr = Presentation(path)
    slides = []
    for s in pr.slides:
        shapes = []
        for sh in s.shapes:
            shapes.append({'type': str(sh.shape_type), 'text': sh.text_frame.text if sh.has_text_frame else None, 'table': [[c.text for c in r.cells] for r in sh.table.rows] if sh.has_table else None})
        slides.append({'shapes': shapes, 'notes': s.notes_slide.notes_text_frame.text if s.has_notes_slide else ''})
    out['slides'] = slides
print(json.dumps(out, ensure_ascii=False, default=str))
